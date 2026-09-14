import json, os
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE=os.path.dirname(os.path.abspath(__file__))
OUT=os.path.join(BASE,'Bella_Car_Search.xlsx')
LAY=json.load(open(os.path.join(BASE,'layout.json')))
H=LAY['H']; FIRST=LAY['first']; LAST=LAY['last']
FONT='Arial'; NAVY='1F3864'
thin=Side(style='thin',color='BFBFBF'); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
wb=load_workbook(OUT)
sh=wb.create_sheet('Start Here',0)
AC="'All Cars'"
def L(n): return get_column_letter(H[n])
MEETS=f"{AC}!${L('Meets All')}${FIRST}:${L('Meets All')}${LAST}"
CPO=f"{AC}!${L('CPO')}${FIRST}:${L('CPO')}${LAST}"

sh.column_dimensions['A'].width=3
sh.column_dimensions['B'].width=34
sh.column_dimensions['C'].width=84
sh.column_dimensions['D'].width=14

def h1(r,t):
    c=sh.cell(row=r,column=2,value=t); c.font=Font(name=FONT,size=16,bold=True,color=NAVY)
def h2(r,t):
    c=sh.cell(row=r,column=2,value=t); c.font=Font(name=FONT,size=12,bold=True,color='2E5C8A')
    sh.cell(row=r,column=2).border=Border(bottom=Side(style='medium',color='2E5C8A'))
    sh.cell(row=r,column=3).border=Border(bottom=Side(style='medium',color='2E5C8A'))
def kv(r,k,v,bold=False,fmt=None):
    a=sh.cell(row=r,column=2,value=k); a.font=Font(name=FONT,size=10,bold=True)
    a.alignment=Alignment(vertical='top')
    b=sh.cell(row=r,column=3,value=v); b.font=Font(name=FONT,size=10,bold=bold)
    b.alignment=Alignment(wrap_text=True,vertical='top')
    if fmt: b.number_format=fmt
    return b
def note(r,t,size=9):
    c=sh.cell(row=r,column=2,value=t); c.font=Font(name=FONT,size=size,italic=True,color='595959')
    sh.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    c.alignment=Alignment(wrap_text=True,vertical='top')

h1(2,"Bella's Car Search")
sh['B3']='Live CARFAX inventory - 2022+ Toyota / Honda / Subaru SUVs within 75 miles of Elkhart, IN'
sh['B3'].font=Font(name=FONT,size=11,italic=True,color='595959')
kv(4,'Data pulled','September 14, 2026 from CARFAX used-car listings')
kv(5,'Search area','Elkhart, IN 46514 - 75 mile radius (South Bend, Mishawaka, Goshen, Warsaw, Fort Wayne, Kalamazoo, Niles)')

h2(7,'The Funnel')
sh['B8']='Filter applied'; sh['C8']='Cars remaining'
for c in (2,3):
    x=sh.cell(row=8,column=c); x.font=Font(name=FONT,size=10,bold=True,color='FFFFFF')
    x.fill=PatternFill('solid',fgColor=NAVY); x.alignment=Alignment(horizontal='left'); x.border=BOX
funnel=[('All 2022+ RAV4 / 4Runner / CR-V / HR-V / Outback under $40k in radius',262),
 ('...that are AWD or 4WD',234),('...with no accidents reported',173),('...with service history on file',168),
 ('...under your mileage cap (Toyota 100k / Honda 70k / Subaru 50k)',144),
 ('...rated GOOD or GREAT value by CARFAX',53),('...with heated seats',36),
 ('...with Apple CarPlay (standard on all of these - no cars lost)',36),
 ('...that are also Certified Pre-Owned',6)]
for i,(t,n) in enumerate(funnel):
    R=9+i
    a=sh.cell(row=R,column=2,value=t); a.font=Font(name=FONT,size=10); a.border=BOX
    b=sh.cell(row=R,column=3,value=n); b.font=Font(name=FONT,size=10,bold=(i in (6,8))); b.border=BOX
    b.alignment=Alignment(horizontal='left')
    if i==6: b.fill=PatternFill('solid',fgColor='C6EFCE')
    if i==8: b.fill=PatternFill('solid',fgColor='FFF2CC')
note(18,'Certified Pre-Owned is scored as a strong bonus rather than a hard filter, so the Shortlist holds 36 cars and the 6 certified ones float near the top. To see only CPO, set "Require Certified Pre-Owned" to TRUE on the Search Criteria tab.')

h2(20,'The Tabs')
tabs=[('Shortlist','The 36 cars that meet every one of your criteria, ranked best first. Start here.'),
 ('Top 10 Compare','The ten best side by side - price, miles, history, features, dealer - for a quick head-to-head.'),
 ('All Cars','All 262 listings with every data point. Use the filter arrows to explore or loosen a requirement.'),
 ('Search Criteria','The yellow cells that drive the YES/no "Meets All" column. Change a cap and the whole sheet re-flags.'),
 ('Scoring Weights','The yellow cells that drive the Score. Change a weight and every car instantly re-ranks.'),
 ('Model Benchmarks','What each model-year actually sells for here, so you can tell a real deal from a normal price.')]
for i,(t,d) in enumerate(tabs):
    R=21+i
    a=sh.cell(row=R,column=2,value=t); a.font=Font(name=FONT,size=10,bold=True,color='2E5C8A'); a.border=BOX
    b=sh.cell(row=R,column=3,value=d); b.font=Font(name=FONT,size=10); b.border=BOX
    b.alignment=Alignment(wrap_text=True,vertical='top')

h2(28,'How the Score Works')
note(29,'Every car gets six sub-scores from 0 to 100, blended using the weights you set. The weighting follows the priority you gave: reliability first, then price, then mileage, then features.')
sc=[('Reliability & History','26%',"CARFAX's model reliability rating, projected repair cost and risk (60%), blended with this specific car's history - accident-free, one owner, personal use, how many service records (40%)."),
 ('Price / Value','22%','The CARFAX Great/Good/Fair value badge (55%), blended with how far below the going rate it is priced (45%). The going rate is fitted from the 262 real listings in this sheet, adjusting for year and mileage.'),
 ('Mileage','18%','Miles against your cap for that make (70%), plus miles-per-year against a 12,000/yr normal (30%).'),
 ('Certified Pre-Owned','15%','100 if factory certified, 0 if not.'),
 ('Features','14%','Heated steering, sunroof, blind spot, power liftgate, leather, remote start, navigation, wireless charging, parking sensors, power seat, premium audio.'),
 ('Distance','5%','100 at your doorstep, 0 at 75 miles. Tiebreaker only.')]
sh['B30']='Component'; sh['C30']='What it measures'; sh['D30']='Weight'
for c in (2,3,4):
    x=sh.cell(row=30,column=c); x.font=Font(name=FONT,size=10,bold=True,color='FFFFFF')
    x.fill=PatternFill('solid',fgColor=NAVY); x.border=BOX; x.alignment=Alignment(horizontal='left')
for i,(n,w,d) in enumerate(sc):
    R=31+i
    a=sh.cell(row=R,column=2,value=n); a.font=Font(name=FONT,size=10,bold=True); a.border=BOX
    b=sh.cell(row=R,column=3,value=d); b.font=Font(name=FONT,size=10); b.border=BOX
    b.alignment=Alignment(wrap_text=True,vertical='top'); sh.row_dimensions[R].height=28
    c=sh.cell(row=R,column=4,value=f"='Scoring Weights'!B{5+i}/100"); c.number_format='0%'
    c.font=Font(name=FONT,size=10,bold=True); c.border=BOX; c.alignment=Alignment(horizontal='center')

h2(39,'Live Counts')
kv(40,'Cars meeting every criterion',f'=COUNTIF({MEETS},"YES")',bold=True,fmt='0')
kv(41,'Of those, Certified Pre-Owned',f'=COUNTIFS({MEETS},"YES",{CPO},"Yes")',bold=True,fmt='0')
kv(42,'Total listings tracked',f'=COUNTA({AC}!${L("VIN")}${FIRST}:${L("VIN")}${LAST})',bold=True,fmt='0')
note(43,'These recalculate from the Search Criteria tab, so you can see instantly how many cars survive a change before you go looking.')

h2(45,'Before You Buy - Read This')
cav=[('Verify CPO with the dealer','CARFAX flags certification from the dealer feed. Confirm it is manufacturer certified (Toyota TCUV, Honda HondaTrue, Subaru Certified) and not a dealer in-house warranty - the warranty difference is significant.'),
 ('Features come from the trim, not the VIN','Heated seats, sunroof and the rest are read from the factory equipment list for that trim. It is accurate for standard equipment but can miss optional packages. Confirm on the window sticker.'),
 ('"Est. Market" is our estimate, not CARFAX','It is a fit across the 262 listings here, adjusting for year and mileage within each model. Treat it as a sanity check on the asking price, not an appraisal.'),
 ('Prices move','Listings change daily and the good ones go fast. The "Days Listed" and "Price Drop" columns tell you which have been sitting - those have the most negotiating room.'),
 ('Get a pre-purchase inspection','Even on a certified car with clean history, an independent inspection is the best $150 you will spend.'),
 ('Mileage caps are yours, not the market','Your Subaru cap of 50k is tight and cuts a lot of Outbacks. If you like the Outback, loosening that cap on the Search Criteria tab opens up more.')]
for i,(t,d) in enumerate(cav):
    R=46+i
    a=sh.cell(row=R,column=2,value=t); a.font=Font(name=FONT,size=10,bold=True); a.border=BOX
    a.alignment=Alignment(vertical='top',wrap_text=True)
    b=sh.cell(row=R,column=3,value=d); b.font=Font(name=FONT,size=10); b.border=BOX
    b.alignment=Alignment(wrap_text=True,vertical='top'); sh.row_dimensions[R].height=34
sh.sheet_view.showGridLines=False
wb.save(OUT)
print('start here added')
