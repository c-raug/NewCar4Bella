#!/usr/bin/env python3
"""Recalculate an .xlsx with LibreOffice so formula results are cached in the file.

openpyxl writes formulas with no cached values, so anything that reads values
rather than formulas (pandas, previewers, GitHub) sees blanks until this runs.
Excel and Google Sheets recalculate on open regardless, so this is about making
the file readable everywhere.

Usage: python3 scripts/recalc.py output/Bella_Car_Search.xlsx [timeout_seconds]
Requires: libreoffice-calc  (apt-get install -y libreoffice-calc)
"""
import json, os, shutil, subprocess, sys, tempfile

def main():
    if len(sys.argv) < 2:
        print(json.dumps({"error": "usage: recalc.py <file.xlsx> [timeout]"})); return 1
    path = os.path.abspath(sys.argv[1])
    timeout = int(sys.argv[2]) if len(sys.argv) > 2 else 300
    if not os.path.exists(path):
        print(json.dumps({"error": f"not found: {path}"})); return 1
    if not shutil.which("soffice"):
        print(json.dumps({"error": "soffice not found; apt-get install -y libreoffice-calc"})); return 1

    with tempfile.TemporaryDirectory() as tmp:
        profile, outdir = os.path.join(tmp, "profile"), os.path.join(tmp, "out")
        os.makedirs(outdir, exist_ok=True)
        cmd = ["soffice", "--headless", "--norestore", "--invisible", "--nolockcheck",
               "--nodefault", "--nofirststartwizard",
               f"-env:UserInstallation=file://{profile}",
               "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", outdir, path]
        try:
            p = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
        except subprocess.TimeoutExpired:
            print(json.dumps({"error": f"LibreOffice timed out after {timeout}s"})); return 1
        produced = os.path.join(outdir, os.path.basename(path))
        if not os.path.exists(produced):
            print(json.dumps({"error": "conversion produced no file",
                              "stdout": p.stdout[-500:], "stderr": p.stderr[-500:]})); return 1
        shutil.copyfile(produced, path)

    try:
        from openpyxl import load_workbook
    except ImportError:
        print(json.dumps({"status": "success", "note": "openpyxl absent; errors not checked"})); return 0
    XL_ERRORS = {"#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A", "#NULL!",
                 "#NUM!", "#GETTING_DATA", "#SPILL!", "#CALC!"}
    wb = load_workbook(path, data_only=True)
    errs, n = {}, 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip() in XL_ERRORS:
                    n += 1
                    errs.setdefault(c.value, []).append(f"{ws.title}!{c.coordinate}")
    print(json.dumps({"status": "errors_found" if n else "success", "total_errors": n,
                      "error_summary": {k: v[:20] for k, v in errs.items()}}, indent=2))
    return 0

if __name__ == "__main__":
    sys.exit(main())
